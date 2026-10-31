"""Perakit modul Imunologi (Markdown -> DOCX).

Pemakaian:
    python build.py            # modul lengkap + DOCX per BAB
Lalu jalankan isi_daftar_isi.py untuk mengisi daftar isi dan membuat PDF.

Konvensi naskah (naskah/*.md):
    - Baris pertama: '## BAB n. Judul' (Heading 2). Subjudul '### ...'.
    - Sitasi: [@kunci] atau [@a; @b]; kunci = ID di ingatan/sumber.ris.
    - Gambar: ![Gambar #. Keterangan.[@kunci]](gambar/berkas.png){width=13cm}
    - Tabel: baris '**Tabel #. Judul**' tepat sebelum tabel pipa.
"""
import copy
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AKAR = os.path.dirname(os.path.abspath(__file__))
NASKAH = os.path.join(AKAR, "naskah")
KELUARAN = os.path.join(AKAR, "keluaran")
GAYA = os.path.join(AKAR, "gaya_referensi.docx")
RIS = os.path.join(AKAR, "ingatan", "sumber.ris")

JUDUL = "MODUL IMUNOLOGI"
SUBJUDUL = "Sistem Imun, Mekanisme Pertahanan, dan Korelasi Klinis"
TAHUN = "2026"

# Susunan BAGIAN -> berkas BAB (urutan = urutan modul)
SUSUNAN = [
    ("BAGIAN I", "DASAR SISTEM IMUN", [
        "BAB_01_Pengantar_Sistem_Imun",
        "BAB_02_Organ_dan_Jaringan_Limfoid",
        "BAB_03_Sel_Sel_Sistem_Imun",
        "BAB_04_Sawar_Epitel_dan_Imunitas_Regional",
    ]),
    ("BAGIAN II", "IMUNITAS BAWAAN", [
        "BAB_05_Pengenalan_Bahaya_dan_Imunitas_Bawaan",
        "BAB_06_Sitokin_dan_Komunikasi_Antarsel",
        "BAB_07_Sistem_Komplemen",
        "BAB_08_Inflamasi_dan_Migrasi_Leukosit",
    ]),
    ("BAGIAN III", "IMUNITAS ADAPTIF", [
        "BAB_09_Antigen_dan_Antibodi",
        "BAB_10_MHC_dan_Presentasi_Antigen",
        "BAB_11_Perkembangan_Limfosit_dan_Diversitas_Reseptor",
        "BAB_12_Aktivasi_dan_Fungsi_Sel_T",
        "BAB_13_Aktivasi_Sel_B_dan_Imunitas_Humoral",
        "BAB_14_Memori_Imun_Vaksinasi_dan_Imunitas_Pasif",
        "BAB_15_Respons_Imun_terhadap_Mikroba",
    ]),
    ("BAGIAN IV", "REGULASI DAN GANGGUAN IMUN", [
        "BAB_16_Toleransi_dan_Dasar_Autoimunitas",
        "BAB_17_Hipersensitivitas_Tipe_I_dan_Alergi",
        "BAB_18_Hipersensitivitas_Tipe_II_III_dan_IV",
        "BAB_19_Imunodefisiensi_Primer",
        "BAB_20_Imunodefisiensi_Sekunder_dan_HIV",
        "BAB_21_Imunologi_Transplantasi",
        "BAB_22_Imunitas_Kehamilan_dan_Jaringan_Khusus",
        "BAB_23_Imunologi_Tumor",
    ]),
    ("BAGIAN V", "PENYAKIT IMUN DAN PENERAPAN", [
        "BAB_24_Autoimunitas_Spesifik_Organ",
        "BAB_25_Autoimunitas_Sistemik",
        "BAB_26_Demam_Rematik",
        "BAB_27_Prinsip_Pemeriksaan_Imunologi",
        "BAB_28_Dasar_Farmakologi_Imun",
        "BAB_29_Psikoneuroimunologi",
        "BAB_30_Integrasi_Kasus",
    ]),
]


# ---------------------------------------------------------------- pustaka
def baca_ris(path):
    rec, data = {}, {}
    for baris in open(path, encoding="utf-8"):
        baris = baris.rstrip("\n")
        m = re.match(r"^([A-Z][A-Z0-9])  - ?(.*)$", baris)
        if not m:
            continue
        tag, isi = m.group(1), m.group(2).strip()
        if tag == "ER":
            data[rec["ID"][0]] = rec
            rec = {}
            continue
        rec.setdefault(tag, []).append(isi)
    return data


def vancouver(r):
    def f(t):
        return r.get(t, [""])[0]
    au = r.get("AU", [])
    nama = []
    for a in au:
        if "," in a:
            sur, ini = [x.strip() for x in a.split(",", 1)]
            ini = "".join(w[0] for w in re.split(r"[\s\.\-]+", ini) if w) if len(ini) > 3 else ini
            nama.append(f"{sur} {ini}")
        else:
            nama.append(a)
    if len(nama) > 6:
        nama = nama[:6]
        dkk = True
    else:
        dkk = bool(r.get("C2"))
    penulis = ", ".join(nama) + (", dkk" if dkk else "")
    if f("TY") == "BOOK":
        s = (penulis + ". " if penulis else "") + f("TI") + ". " + f("ET") + "."
        if f("PB"):
            s += " " + f("PB") + (";" + " " + f("PY") if f("PY") else "") + "."
        if f("N1") and "Bab" in f("N1"):
            s += " " + f("N1") + "."
        return s
    hal = f("SP") + ("-" + f("EP") if f("EP") else "")
    s = f"{penulis}. {f('TI')}. {f('JO')}. {f('PY')};{f('VL')}"
    if f("IS"):
        s += f"({f('IS')})"
    s += f":{hal}."
    if f("DO"):
        s += f" doi:{f('DO')}."
    if f("AN"):
        s += f" {f('AN')}."
    return s


# ---------------------------------------------------------------- naskah
POLA_SITASI = re.compile(r"\[(@[^\]]+)\]")


def baca_naskah(nama):
    return open(os.path.join(NASKAH, nama + ".md"), encoding="utf-8").read()


def nomori_gambar_tabel(teks_list):
    """Ganti 'Gambar #.' dan 'Tabel #.' berurutan di seluruh modul."""
    g = t = 0
    hasil = []
    for teks in teks_list:
        def rg(m):
            nonlocal g
            g += 1
            return f"Gambar {g}."
        def rt(m):
            nonlocal t
            t += 1
            return f"Tabel {t}."
        teks = re.sub(r"Gambar #\.", rg, teks)
        teks = re.sub(r"Tabel #\.", rt, teks)
        hasil.append(teks)
    return hasil


def ganti_sitasi(teks, urutan):
    def r(m):
        kunci = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        nomor = []
        for k in kunci:
            if k not in urutan:
                urutan.append(k)
            nomor.append(urutan.index(k) + 1)
        nomor = sorted(set(nomor))
        return "^" + ",".join(str(n) for n in nomor) + "^"
    return POLA_SITASI.sub(r, teks)


def daftar_pustaka_md(urutan, ris, judul_heading):
    baris = [judul_heading, ""]
    for i, k in enumerate(urutan, 1):
        if k not in ris:
            sys.exit(f"Kunci pustaka tidak ditemukan: {k}")
        baris.append(f"{i}. {vancouver(ris[k])}")
        baris.append("")
    return "\n".join(baris)


# ---------------------------------------------------------------- gaya
def siapkan_gaya():
    if os.path.exists(GAYA):
        return
    tmp = os.path.join(AKAR, "_pandoc_ref.docx")
    subprocess.run(["pandoc", "-o", tmp, "--print-default-data-file", "reference.docx"], check=True)
    d = Document(tmp)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for s in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, s, Cm(2.54))
    st = d.styles
    for nama in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text",
                 "Image Caption", "Table Caption", "Bibliography"):
        try:
            s = st[nama]
        except KeyError:
            continue
        s.font.name = "Times New Roman"
        s.font.size = Pt(12)
        s.font.color.rgb = RGBColor(0, 0, 0)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
        if nama in ("Body Text", "First Paragraph", "Compact", "Normal"):
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    st["Body Text"].paragraph_format.space_after = Pt(6)
    st["First Paragraph"].paragraph_format.space_after = Pt(6)
    st["Compact"].paragraph_format.space_after = Pt(2)
    st["Image Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Image Caption"].font.italic = False
    for nama, ukuran, tebal, miring, garis in (("Heading 1", 16, True, False, False),
                                               ("Heading 2", 17, True, False, False),
                                               ("Heading 3", 13, False, True, True),
                                               ("Title", 20, True, False, False)):
        s = st[nama]
        s.font.name = "Times New Roman"
        s.font.size = Pt(ukuran)
        s.font.bold = tebal
        s.font.italic = miring
        s.font.underline = garis
        s.font.color.rgb = RGBColor(0, 0, 0)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
        for t in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(t)) is not None:
                del rf.attrib[qn(t)]
    st["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Heading 1"].paragraph_format.page_break_before = True
    st["Heading 1"].paragraph_format.space_after = Pt(18)
    h2 = st["Heading 2"]
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(10)
    ppr = h2.element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "18"); b.set(qn("w:space"), "4"); b.set(qn("w:color"), "FFD20A")
    bdr.append(b)
    ppr.append(bdr)
    st["Heading 3"].paragraph_format.space_before = Pt(12)
    st["Heading 3"].paragraph_format.space_after = Pt(4)
    d.save(GAYA)
    os.remove(tmp)


# ---------------------------------------------------------------- pascaproses
def paragraf_baru_sebelum(par, teks="", tebal=False, ukuran=12, miring=False, rata=WD_ALIGN_PARAGRAPH.CENTER,
                          spasi_setelah=0):
    p = par.insert_paragraph_before()
    p.style = par.part.document.styles["Normal"]
    p.alignment = rata
    p.paragraph_format.space_after = Pt(spasi_setelah)
    if teks:
        r = p.add_run(teks)
        r.bold, r.italic = tebal, miring
        r.font.size = Pt(ukuran)
        r.font.name = "Times New Roman"
    return p


def tambah_field(paragraf, instr):
    r = paragraf.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Klik kanan > Update Field untuk memuat daftar isi."
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r._r.append(a); r._r.append(i); r._r.append(s); r._r.append(t); r._r.append(e)


def pecah_halaman_sebelum(par):
    p = par.insert_paragraph_before()
    p.add_run().add_break(WD_BREAK.PAGE)


def rapikan(path, mode, bagian=None):
    d = Document(path)
    body = d.paragraphs
    # gambar dan keterangan rata tengah
    for p in d.paragraphs:
        if p._p.xpath(".//w:drawing") or p.style.name in ("Captioned Figure", "Figure"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name == "Image Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        teks = p.text.strip()
        if re.match(r"^Tabel \d+\.", teks):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
    # tabel: garis atas-bawah, garis tipis antarbaris
    for tb in d.tables:
        tpr = tb._tbl.tblPr
        for lama in tpr.findall(qn("w:tblBorders")):
            tpr.remove(lama)
        bd = OxmlElement("w:tblBorders")
        for nama, sz in (("top", "12"), ("bottom", "12"), ("insideH", "4"),
                         ("left", None), ("right", None), ("insideV", None)):
            e = OxmlElement(f"w:{nama}")
            if sz:
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz); e.set(qn("w:color"), "000000")
            else:
                e.set(qn("w:val"), "nil")
            bd.append(e)
        tpr.append(bd)
        lebar = tpr.find(qn("w:tblW"))
        if lebar is None:
            lebar = OxmlElement("w:tblW"); tpr.append(lebar)
        lebar.set(qn("w:type"), "pct"); lebar.set(qn("w:w"), "5000")
        for i, row in enumerate(tb.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(11)
                        if i == 0:
                            r.bold = True
    # nomor halaman tengah di footer
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for s in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, s, Cm(2.54))
        f = sec.footer
        p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        r._r.append(a); r._r.append(i); r._r.append(e)
        r.font.name = "Times New Roman"; r.font.size = Pt(11)
    pertama = d.paragraphs[0]
    if mode == "lengkap":
        # sampul
        for _ in range(8):
            paragraf_baru_sebelum(pertama)
        paragraf_baru_sebelum(pertama, JUDUL, tebal=True, ukuran=22, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, SUBJUDUL, tebal=True, ukuran=15, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, "Bahan Belajar Tahap Preklinik", miring=True, ukuran=12)
        for _ in range(6):
            paragraf_baru_sebelum(pertama)
        paragraf_baru_sebelum(pertama, "Dokter Edukasi Academy", tebal=True, ukuran=14, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, "Pusat Pembelajaran & Riset Medis", ukuran=12, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, TAHUN, ukuran=12)
        pecah_halaman_sebelum(pertama)
        paragraf_baru_sebelum(pertama, "DAFTAR ISI", tebal=True, ukuran=16, spasi_setelah=12)
        toc = paragraf_baru_sebelum(pertama, rata=WD_ALIGN_PARAGRAPH.LEFT)
        tambah_field(toc, 'TOC \\o "1-2" \\h \\z \\u')
    else:
        paragraf_baru_sebelum(pertama, JUDUL, tebal=True, ukuran=14, spasi_setelah=2)
        paragraf_baru_sebelum(pertama, "Dokter Edukasi Academy · Pusat Pembelajaran & Riset Medis",
                              ukuran=11, spasi_setelah=8)
        if bagian:
            paragraf_baru_sebelum(pertama, bagian, tebal=True, ukuran=12, spasi_setelah=4)
    d.save(path)


def pandoc(md, out):
    tmp = out + ".md"
    open(tmp, "w", encoding="utf-8").write(md)
    subprocess.run(["pandoc", tmp, "-f", "markdown+superscript+pipe_tables+implicit_figures",
                    "-t", "docx", "--reference-doc", GAYA, "--resource-path", AKAR, "-o", out],
                   check=True, cwd=AKAR)
    os.remove(tmp)


def main():
    siapkan_gaya()
    ris = baca_ris(RIS)
    os.makedirs(KELUARAN, exist_ok=True)
    os.makedirs(os.path.join(KELUARAN, "DOCX per BAB"), exist_ok=True)
    daftar = [(bg, judul, nama) for bg, judul, babs in SUSUNAN for nama in babs]
    tersedia = [x for x in daftar if os.path.exists(os.path.join(NASKAH, x[2] + ".md"))]
    teks = nomori_gambar_tabel([baca_naskah(x[2]) for x in tersedia])

    # modul lengkap
    urutan, potong, bagian_lalu = [], [], None
    for (bg, judul, nama), isi in zip(tersedia, teks):
        if bg != bagian_lalu:
            potong.append(f"# {bg}. {judul}\n")
            bagian_lalu = bg
        potong.append(ganti_sitasi(isi, urutan) + "\n")
    potong.append(daftar_pustaka_md(urutan, ris, "# DAFTAR PUSTAKA"))
    out = os.path.join(KELUARAN, "Modul_Imunologi.docx")
    pandoc("\n".join(potong), out)
    rapikan(out, "lengkap")
    print("modul lengkap:", len(tersedia), "BAB,", len(urutan), "pustaka")

    # per BAB
    for (bg, judul, nama), isi in zip(tersedia, teks):
        urutan = []
        md = ganti_sitasi(isi, urutan) + "\n\n" + daftar_pustaka_md(urutan, ris, "### Daftar Pustaka")
        o = os.path.join(KELUARAN, "DOCX per BAB", nama + ".docx")
        pandoc(md, o)
        rapikan(o, "bab", f"{bg}. {judul}")
    print("selesai")


if __name__ == "__main__":
    main()

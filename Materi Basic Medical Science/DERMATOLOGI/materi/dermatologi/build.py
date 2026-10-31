"""Perakit modul Dermatologi (Markdown -> DOCX).

Pemakaian:
    python build.py            # modul lengkap + DOCX per BAB
Lalu jalankan isi_daftar_isi.py untuk mengisi daftar isi dan membuat PDF.

Konvensi naskah (naskah/*.md):
    - Baris pertama: '## BAB n. Judul' (Heading 2). Subjudul '### ...'.
    - Sitasi: [@kunci] atau [@a; @b]; kunci = ID di ingatan/sumber.ris.
    - Gambar: ![Gambar #. Keterangan.[@kunci]](gambar/berkas.png){width=13cm}
    - Tabel: baris '**Tabel #. Judul**' tepat sebelum tabel pipa.
"""
import copy
import json
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AKAR = os.path.dirname(os.path.abspath(__file__))
NASKAH = os.path.join(AKAR, "naskah")
KELUARAN = os.path.join(AKAR, "keluaran")
GAYA = os.path.join(AKAR, "gaya_referensi.docx")
RIS = os.path.join(AKAR, "ingatan", "sumber.ris")

JUDUL = "MODUL DERMATOLOGI"
SUBJUDUL = "Struktur Fungsi dan Korelasi Klinis Kulit"
TAHUN = "2026"

# Susunan BAGIAN -> berkas BAB (urutan = urutan modul)
MANIFEST = json.load(open(os.path.join(AKAR, "ingatan", "manifest.json"), encoding="utf-8"))
NOMOR = json.load(open(os.path.join(AKAR, "ingatan", "nomor_sitasi.json"), encoding="utf-8"))
SUSUNAN = [
    ("BAGIAN I", "STRUKTUR DAN FUNGSI KULIT", [x["nama"] for x in MANIFEST[:11]]),
    ("BAGIAN II", "INFLAMASI DAN INFEKSI KULIT", [x["nama"] for x in MANIFEST[11:23]]),
    ("BAGIAN III", "LESI KULIT DAN KORELASI SISTEMIK", [x["nama"] for x in MANIFEST[23:29]]),
    ("BAGIAN IV", "FARMAKOLOGI PROSEDUR DAN INTEGRASI", [x["nama"] for x in MANIFEST[29:]])
]


# ---------------------------------------------------------------- pustaka
def baca_ris(path):
    rec, data = {}, {}
    for baris in open(path, encoding="utf-8"):
        baris = baris.rstrip("\n")
        m = re.match(r"^([A-Z][A-Z0-9])  - ?(.*)$", baris)
        if not m:
            continue
        tag, isi = m.group(1), m.group(2).strip()
        if tag == "ER":
            data[rec["ID"][0]] = rec
            rec = {}
            continue
        rec.setdefault(tag, []).append(isi)
    return data


def vancouver(r):
    def f(k): return r.get(k,[""])[0]
    aus = r.get("AU",[])
    names=[]
    for a in aus[:6]:
        if "," in a:
            sur, given = a.split(",",1)
            initials = "".join(w[0] for w in given.strip().split()) if len(given.strip())>4 else given.strip()
            names.append(sur+" "+initials)
        else: names.append(a)
    author=", ".join(names)+( ", et al" if len(aus)>6 else "")
    if f("TY")=="ELEC":
        return author+". "+f("TI")+" [Internet]. Diakses 31 Oktober 2026. "+f("UR")+"."
    if f("TY")=="BOOK":
        return author+", editors. "+f("TI")+". "+f("ET")+". "+f("PB")+"; "+f("PY")+". ISBN: "+f("SN")+"."
    pages=f("SP")+("-"+f("EP") if f("EP") else "")
    doi=" https://doi.org/"+f("DO") if f("DO") else ""
    return author+". "+f("TI")+". "+f("JO")+". "+f("PY")+";"+f("VL")+("("+f("IS")+")" if f("IS") else "")+(":"+pages if pages else "")+"."+doi+". "+f("AN")+"."


# ---------------------------------------------------------------- naskah
POLA_SITASI = re.compile(r"\[(@[^\]]+)\]")


def baca_naskah(nama):
    teks = open(os.path.join(NASKAH, nama + ".md"), encoding="utf-8").read()
    def ubah(m):
        return "[" + "; ".join("@" + NOMOR[int(n)-1] for n in m.group(1).split(",")) + "]"
    return re.sub(r"\^([0-9,]+)\^", ubah, teks)


def nomori_gambar_tabel(teks_list):
    """Ganti 'Gambar #.' dan 'Tabel #.' berurutan di seluruh modul."""
    g = t = 0
    hasil = []
    for teks in teks_list:
        def rg(m):
            nonlocal g
            g += 1
            return f"Gambar {g}."
        def rt(m):
            nonlocal t
            t += 1
            return f"Tabel {t}."
        teks = re.sub(r"Gambar #\.", rg, teks)
        teks = re.sub(r"Tabel #\.", rt, teks)
        hasil.append(teks)
    return hasil


def ganti_sitasi(teks, urutan):
    def r(m):
        kunci = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        nomor = []
        for k in kunci:
            if k not in urutan:
                urutan.append(k)
            nomor.append(urutan.index(k) + 1)
        nomor = sorted(set(nomor))
        return "^" + ",".join(str(n) for n in nomor) + "^"
    return POLA_SITASI.sub(r, teks)


def daftar_pustaka_md(urutan, ris, judul_heading):
    baris = [judul_heading, ""]
    for i, k in enumerate(urutan, 1):
        if k not in ris:
            sys.exit(f"Kunci pustaka tidak ditemukan: {k}")
        baris.append(f"{i}. {vancouver(ris[k])}")
        baris.append("")
    return "\n".join(baris)


# ---------------------------------------------------------------- gaya
def siapkan_gaya():
    if os.path.exists(GAYA):
        return
    tmp = os.path.join(AKAR, "_pandoc_ref.docx")
    subprocess.run(["pandoc", "-o", tmp, "--print-default-data-file", "reference.docx"], check=True)
    d = Document(tmp)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for s in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, s, Cm(2.54))
    st = d.styles
    for nama in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text",
                 "Image Caption", "Table Caption", "Bibliography"):
        try:
            s = st[nama]
        except KeyError:
            continue
        s.font.name = "Times New Roman"
        s.font.size = Pt(12)
        s.font.color.rgb = RGBColor(0, 0, 0)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
        if nama in ("Body Text", "First Paragraph", "Compact", "Normal"):
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    st["Body Text"].paragraph_format.space_after = Pt(6)
    st["First Paragraph"].paragraph_format.space_after = Pt(6)
    st["Compact"].paragraph_format.space_after = Pt(2)
    st["Image Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Image Caption"].font.italic = False
    for nama, ukuran, tebal, miring, garis in (("Heading 1", 16, True, False, False),
                                               ("Heading 2", 17, True, False, False),
                                               ("Heading 3", 13, False, True, True),
                                               ("Title", 20, True, False, False)):
        s = st[nama]
        s.font.name = "Times New Roman"
        s.font.size = Pt(ukuran)
        s.font.bold = tebal
        s.font.italic = miring
        s.font.underline = garis
        s.font.color.rgb = RGBColor(0, 0, 0)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
        for t in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(t)) is not None:
                del rf.attrib[qn(t)]
    st["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Heading 1"].paragraph_format.page_break_before = True
    st["Heading 1"].paragraph_format.space_after = Pt(18)
    h2 = st["Heading 2"]
    h2.paragraph_format.page_break_before = True
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(10)
    ppr = h2.element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "18"); b.set(qn("w:space"), "4"); b.set(qn("w:color"), "FFD20A")
    bdr.append(b)
    ppr.append(bdr)
    st["Heading 3"].paragraph_format.space_before = Pt(12)
    st["Heading 3"].paragraph_format.space_after = Pt(4)
    d.save(GAYA)
    os.remove(tmp)


# ---------------------------------------------------------------- pascaproses
def paragraf_baru_sebelum(par, teks="", tebal=False, ukuran=12, miring=False, rata=WD_ALIGN_PARAGRAPH.CENTER,
                          spasi_setelah=0):
    p = par.insert_paragraph_before()
    p.style = par.part.document.styles["Normal"]
    p.alignment = rata
    p.paragraph_format.space_after = Pt(spasi_setelah)
    if teks:
        if teks == JUDUL: p.style = par.part.document.styles["Title"]
        r = p.add_run(teks)
        r.bold, r.italic = tebal, miring
        r.font.size = Pt(ukuran)
        r.font.name = "Times New Roman"
    return p


def tambah_field(paragraf, instr):
    r = paragraf.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Klik kanan > Update Field untuk memuat daftar isi."
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r._r.append(a); r._r.append(i); r._r.append(s); r._r.append(t); r._r.append(e)


def pecah_halaman_sebelum(par):
    p = par.insert_paragraph_before()
    p.add_run().add_break(WD_BREAK.PAGE)


def rapikan(path, mode, bagian=None):
    d = Document(path)
    body = d.paragraphs
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in d.styles:
            st=d.styles[name]
            st.font.name="Times New Roman"; st.font.size=Pt(12)
            st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    d.styles["Heading 1"].paragraph_format.page_break_before=True
    d.styles["Heading 2"].paragraph_format.page_break_before=True
    d.styles["Heading 2"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    prev_heading=None
    for par in d.paragraphs:
        if par.style.name=="Heading 2":
            par.paragraph_format.page_break_before = mode=="lengkap" and prev_heading!="Heading 1"
        if par.style.name in ("Heading 1", "Heading 2"): prev_heading=par.style.name
    # gambar dan keterangan rata tengah
    for p in d.paragraphs:
        if p._p.xpath(".//w:drawing") or p.style.name in ("Captioned Figure", "Figure"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name == "Image Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        teks = p.text.strip()
        if re.match(r"^Tabel \d+\.", teks):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
    # tabel: garis atas-bawah, garis tipis antarbaris
    for tb in d.tables:
        tpr = tb._tbl.tblPr
        for lama in tpr.findall(qn("w:tblBorders")):
            tpr.remove(lama)
        bd = OxmlElement("w:tblBorders")
        for nama, sz in (("top", "12"), ("bottom", "12"), ("insideH", "4"),
                         ("left", None), ("right", None), ("insideV", None)):
            e = OxmlElement(f"w:{nama}")
            if sz:
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz); e.set(qn("w:color"), "000000")
            else:
                e.set(qn("w:val"), "nil")
            bd.append(e)
        tpr.append(bd)
        lebar = tpr.find(qn("w:tblW"))
        if lebar is None:
            lebar = OxmlElement("w:tblW"); tpr.append(lebar)
        lebar.set(qn("w:type"), "pct"); lebar.set(qn("w:w"), "5000")
        for i, row in enumerate(tb.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(12)
                        if i == 0:
                            r.bold = True
    # nomor halaman tengah di footer
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for s in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, s, Cm(2.54))
        f = sec.footer
        p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        r._r.append(a); r._r.append(i); r._r.append(e)
        r.font.name = "Times New Roman"; r.font.size = Pt(12)
    pertama = d.paragraphs[0]
    if mode == "lengkap":
        # sampul
        for _ in range(8):
            paragraf_baru_sebelum(pertama)
        paragraf_baru_sebelum(pertama, JUDUL, tebal=True, ukuran=22, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, SUBJUDUL, tebal=True, ukuran=15, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, "Bahan Belajar Tahap Preklinik", miring=True, ukuran=12)
        for _ in range(6):
            paragraf_baru_sebelum(pertama)
        paragraf_baru_sebelum(pertama, "Dokter Edukasi Academy", tebal=True, ukuran=14, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, "Pusat Pembelajaran & Riset Medis", ukuran=12, spasi_setelah=6)
        paragraf_baru_sebelum(pertama, TAHUN, ukuran=12)
        pecah_halaman_sebelum(pertama)
        paragraf_baru_sebelum(pertama, "DAFTAR ISI", tebal=True, ukuran=16, spasi_setelah=12)
        toc = paragraf_baru_sebelum(pertama, rata=WD_ALIGN_PARAGRAPH.LEFT)
        tambah_field(toc, 'TOC \\o "1-2" \\h \\z \\u')
    else:
        paragraf_baru_sebelum(pertama, JUDUL, tebal=True, ukuran=14, spasi_setelah=2)
        paragraf_baru_sebelum(pertama, "Dokter Edukasi Academy · Pusat Pembelajaran & Riset Medis",
                              ukuran=11, spasi_setelah=8)
        if bagian:
            paragraf_baru_sebelum(pertama, bagian, tebal=True, ukuran=12, spasi_setelah=4)
    d.save(path)


def pandoc(md, out):
    tmp = out + ".md"
    open(tmp, "w", encoding="utf-8").write(md)
    subprocess.run(["pandoc", tmp, "-f", "markdown+superscript+pipe_tables+implicit_figures",
                    "-t", "docx", "--reference-doc", GAYA, "--resource-path", AKAR, "-o", out],
                   check=True, cwd=AKAR)
    os.remove(tmp)


def main():
    siapkan_gaya()
    ris = baca_ris(RIS)
    os.makedirs(KELUARAN, exist_ok=True)
    os.makedirs(os.path.join(KELUARAN, "DOCX per BAB"), exist_ok=True)
    daftar = [(bg, judul, nama) for bg, judul, babs in SUSUNAN for nama in babs]
    tersedia = [x for x in daftar if os.path.exists(os.path.join(NASKAH, x[2] + ".md"))]
    teks = nomori_gambar_tabel([baca_naskah(x[2]) for x in tersedia])

    # modul lengkap
    urutan, potong, bagian_lalu = [], [], None
    for (bg, judul, nama), isi in zip(tersedia, teks):
        if bg != bagian_lalu:
            potong.append(f"# {bg}. {judul}\n")
            bagian_lalu = bg
        potong.append(ganti_sitasi(isi, urutan) + "\n")
    potong.append(daftar_pustaka_md(urutan, ris, "# DAFTAR PUSTAKA"))
    out = os.path.join(KELUARAN, "Modul_Dermatologi_Preklinik.docx")
    pandoc("\n".join(potong), out)
    rapikan(out, "lengkap")
    print("modul lengkap:", len(tersedia), "BAB,", len(urutan), "pustaka")

    # per BAB
    for (bg, judul, nama), isi in zip(tersedia, teks):
        urutan = []
        md = ganti_sitasi(isi, urutan) + "\n\n" + daftar_pustaka_md(urutan, ris, "# DAFTAR PUSTAKA")
        o = os.path.join(KELUARAN, "DOCX per BAB", nama + ".docx")
        pandoc(md, o)
        rapikan(o, "bab", f"{bg}. {judul}")
    print("selesai")


if __name__ == "__main__":
    main()

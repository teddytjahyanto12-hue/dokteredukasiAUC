"""Audit struktur dan render semua halaman keluaran untuk inspeksi visual."""
from pathlib import Path
import json
import re
import hashlib
import pymupdf
from PIL import Image, ImageDraw
from docx import Document

B=Path(__file__).resolve().parents[1]
K=B/'keluaran'
Q=B/'qa'

def main():
    issues=[]
    nas=sorted((B/'naskah').glob('BAB_*.md'))
    assert len(nas)==32
    seen=[]
    needed=['Tujuan belajar','Konsep dasar','Gambar','Hubungan sebab-akibat','Korelasi klinis','Rangkuman','Soal aplikasi beserta kunci']
    for n,p in enumerate(nas,1):
        s=p.read_text(encoding='utf-8')
        assert s.startswith(f'## BAB {n}.')
        assert re.findall(r'^### (.+)',s,re.M)==needed,p.name
        for name in ['Konsep dasar','Hubungan sebab-akibat','Korelasi klinis','Rangkuman']:
            section=s.split('### '+name)[1].split('### ')[0]
            assert re.search(r'\^[0-9,]+\^',section),(p.name,name)
        for x in re.findall(r'\^([0-9,]+)\^',s):
            for v in map(int,x.split(',')):
                if v not in seen:seen.append(v)
        assert not re.search(r'catatan\s*siswa|catatansiswa|\bAI\b|—|--|Fitzpatrick|Sherwood|OpenStax|\bPDF\b',re.sub(r'^\|[-| :]+\|$','',s,flags=re.M),re.I),p.name
    assert seen==list(range(1,11)),seen
    manifest=[];records=[]
    files=[K/'Modul_Dermatologi_Preklinik.pdf']+sorted((K/'PDF per BAB').glob('*.pdf'))
    assert len(files)==33
    for fi,p in enumerate(files):
        doc=pymupdf.open(p)
        for j,page in enumerate(doc):
            txt=page.get_text()
            if re.search(r'catatan\s*siswa|catatansiswa|\bAI\b|—|--|Klik kanan|Update Field|\[@',txt,re.I):issues.append([p.name,j+1,'larangan/placeholder'])
            if len(txt.strip())<6:issues.append([p.name,j+1,'halaman kosong'])
            for block in page.get_text('dict')['blocks']:
                x0,y0,x1,y1=block['bbox']
                if x0<50 or x1>page.rect.width-45 or y0<40 or y1>page.rect.height-25:issues.append([p.name,j+1,'batas halaman',block['bbox']])
            png=Q/f'final_{fi:02}_{j+1:03}.png'
            page.get_pixmap(matrix=pymupdf.Matrix(1.1,1.1)).save(str(png))
            records.append([p.name,j+1,png])
        manifest.append({'berkas':str(p.relative_to(K)),'halaman':len(doc),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    for start in range(0,len(records),4):
        canvas=Image.new('RGB',(1300,1860),'#d9dfe3');draw=ImageDraw.Draw(canvas)
        for i,(name,n,png) in enumerate(records[start:start+4]):
            im=Image.open(png);im.thumbnail((640,905));x=(i%2)*650;y=(i//2)*930
            canvas.paste(im,(x,y+25));draw.text((x+4,y+5),f'{name[:45]} | {n}',fill='black')
        canvas.save(Q/f'final_kontak_{start//4+1:02}.jpg',quality=92)
    d=Document(K/'Modul_Dermatologi_Preklinik.docx')
    assert len(d.inline_shapes)==32
    assert len(d.tables)==2
    assert len([p for p in d.paragraphs if p.style.name=='Heading 2'])==32
    assert all(abs(s.page_width.cm-21)<.02 and abs(s.page_height.cm-29.7)<.02 for s in d.sections)
    assert all(abs(s.left_margin.cm-2.54)<.02 for s in d.sections)
    full=pymupdf.open(files[0]);toc=full[1].get_text()
    assert 'BAB 32.' in toc and 'DAFTAR PUSTAKA' in toc
    assert sorted(set(map(int,re.findall(r'Gambar (\d+)\.', '\n'.join(p.get_text() for p in full[2:])))))==list(range(1,33))
    assert not issues,issues
    report={'bab':32,'pustaka':10,'artikel_doi_pmid':8,'gambar':32,'tabel':2,
            'halaman_modul':len(full),'halaman_seluruh_pdf':len(records),'lembar_inspeksi':(len(records)+3)//4,
            'issues':issues,'berkas':manifest}
    (Q/'hasil_final.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='berkas'},ensure_ascii=False))

if __name__=='__main__':main()
